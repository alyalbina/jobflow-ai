from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from jobflow.app import create_app
from jobflow.storage import now


def test_invalid_manual_fields_cannot_break_matching(authenticated):
    client, store, headers = authenticated
    job = store.ingest(1, 812, "Vacancy: Product Analyst\nRequirements: SQL", now())["job_id"]
    for body in ({"required_skills": "SQL"}, {"years_experience": "2"}, {"country_restricted": "false"}):
        assert client.patch(f"/api/jobs/{job}", json=body, headers=headers).status_code == 422
    assert client.get("/api/state").status_code == 200


def test_manual_import_keeps_catchup_cursor_and_source(authenticated):
    client, store, headers = authenticated
    result = client.post(
        "/api/import-message",
        json={
            "source_id": 1,
            "message_id": 200,
            "text": "Vacancy: Product Analyst\nRequirements: SQL",
            "published_at": now(),
            "original_url": "https://t.me/revacancy_global/200",
        },
        headers=headers,
    )
    assert result.status_code == 200
    assert store.rows("SELECT cursor FROM sources WHERE id=1")[0]["cursor"] == 0
    assert store.job(result.json()["job_id"])["original_url"].endswith("/200")


def test_weekly_effort_is_actual_owner_entry_not_a_savings_claim(authenticated):
    client, store, headers = authenticated
    payload = {
        "week": "2026-W41",
        "phase": "baseline",
        "search_minutes": 180,
        "preparation_minutes": 90,
        "qualified_applications": 3,
    }
    assert client.post("/api/effort", json=payload, headers=headers).status_code == 200
    payload["search_minutes"] = 200
    assert client.post("/api/effort", json=payload, headers=headers).status_code == 200
    assert len(store.analytics()["effort_logs"]) == 1
    assert store.analytics()["effort_logs"][0]["search_minutes"] == 200
    payload["search_minutes"] = -1
    assert client.post("/api/effort", json=payload, headers=headers).status_code == 422


def test_tailored_cv_keeps_original_evidence_provenance(authenticated):
    client, store, headers = authenticated
    doc = Document()
    doc.add_paragraph("I built SQL retention analyses for a fictional test project.")
    buffer = BytesIO()
    doc.save(buffer)
    uploaded = client.post(
        "/api/cvs", headers=headers, files={"file": ("test.docx", buffer.getvalue())}
    ).json()
    fact = uploaded["suggested_facts"][0]
    fact.update(status="verified", approved=True)
    assert client.post("/api/profile", json={"facts": [fact]}, headers=headers).status_code == 200
    job = store.ingest(1, 700, "Vacancy: Product Analyst\nRequirements: SQL", now())["job_id"]
    cv = client.post("/api/tailor", json={"job_id": job, "cv_id": uploaded["id"]}, headers=headers).json()
    letter = client.post("/api/letters", json={"job_id": job, "cv_id": cv["id"]}, headers=headers)
    assert letter.status_code == 200
    assert letter.json()["evidence"][0]["cv_id"] == uploaded["id"]


def test_letter_link_cannot_reference_other_vacancy(authenticated):
    client, store, headers = authenticated
    first = store.ingest(1, 800, "Vacancy: Product Analyst\nRequirements: SQL", now())["job_id"]
    second = store.ingest(1, 801, "Vacancy: Business Analyst\nRequirements: API", now())["job_id"]
    client.post(
        "/api/profile",
        json={
            "facts": [
                {
                    "id": "f",
                    "text": "I wrote SQL analyses.",
                    "skills": ["SQL"],
                    "status": "self_reported",
                    "approved": True,
                }
            ]
        },
        headers=headers,
    )
    letter = client.post("/api/letters", json={"job_id": first}, headers=headers).json()
    result = client.post(
        "/api/applications",
        json={"job_id": second, "status": "Preparing", "letter_id": letter["id"]},
        headers=headers,
    )
    assert result.status_code == 422


@pytest.fixture
def authenticated(tmp_path):
    app = create_app(tmp_path, bootstrap="synthetic-test-token", testing=True)
    with TestClient(app) as client:
        response = client.post(
            "/api/bootstrap", headers={"Origin": "http://testserver"}, json={"token": "synthetic-test-token"}
        )
        headers = {"Origin": "http://testserver", "X-JobFlow-CSRF": response.json()["csrf"]}
        yield client, app.state.store, headers


def test_one_time_login_origin_csrf_and_host(tmp_path):
    app = create_app(tmp_path, bootstrap="fake-key", testing=True)
    with TestClient(app) as client:
        assert client.get("/api/state").status_code == 401
        assert (
            client.post(
                "/api/bootstrap", headers={"Origin": "https://evil.example"}, json={"token": "fake-key"}
            ).status_code
            == 403
        )
        r = client.post("/api/bootstrap", headers={"Origin": "http://testserver"}, json={"token": "fake-key"})
        assert r.status_code == 200
        assert "HttpOnly" in r.headers["set-cookie"] and "SameSite=strict" in r.headers["set-cookie"]
        assert (
            client.post(
                "/api/bootstrap", headers={"Origin": "http://testserver"}, json={"token": "fake-key"}
            ).status_code
            == 401
        )
        assert (
            client.post(
                "/api/profile", headers={"Origin": "http://testserver"}, json={"name": "Fake"}
            ).status_code
            == 403
        )
        assert client.get("/api/state", headers={"Host": "evil.example"}).status_code == 400
        assert client.get("/private/jobflow.sqlite").status_code == 404


def test_manual_apply_confirm_and_foreign_keys(authenticated):
    client, store, headers = authenticated
    job = store.ingest(1, 1, "Vacancy: Product Analyst\nRequirements: SQL", now())["job_id"]
    assert (
        client.post(
            "/api/applications", headers=headers, json={"job_id": job, "status": "Preparing"}
        ).status_code
        == 200
    )
    assert (
        client.post(
            "/api/applications", headers=headers, json={"job_id": job, "status": "Applied"}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/applications", headers=headers, json={"job_id": job, "status": "Applied", "confirmed": True}
        ).status_code
        == 200
    )
    state = client.get("/api/state").json()
    assert len(state["events"]) == 2 and state["applications"][0]["applied_at"]


def test_cv_import_verification_export_and_tampered_claim(authenticated):
    client, store, headers = authenticated
    doc = Document()
    doc.add_paragraph("I used SQL for a fictional retention analysis.")
    buffer = BytesIO()
    doc.save(buffer)
    r = client.post(
        "/api/cvs",
        headers=headers,
        files={"file": ("fictional.docx", buffer.getvalue())},
        data={"label": "Synthetic CV"},
    )
    assert r.status_code == 200
    assert r.json()["suggested_facts"][0]["status"] == "inferred"
    fact = r.json()["suggested_facts"][0]
    fact.update(status="verified", approved=True)
    assert client.post("/api/profile", headers=headers, json={"facts": [fact]}).status_code == 200
    fact["skills"] = ["Machine Learning"]
    assert client.post("/api/profile", headers=headers, json={"facts": [fact]}).status_code == 422
    fact["text"] = "I increased revenue by 300%."
    assert client.post("/api/profile", headers=headers, json={"facts": [fact]}).status_code == 422
    exported = client.get(f"/api/cvs/{r.json()['id']}/export/docx")
    assert exported.status_code == 200
    assert Document(BytesIO(exported.content)).paragraphs[0].text.startswith("I used SQL")


def test_corrections_preserve_original_text(authenticated):
    client, store, headers = authenticated
    original = "Vacancy: Product Manager\nRequirements: Discovery"
    job = store.ingest(1, 1, original, now())["job_id"]
    assert (
        client.patch(f"/api/jobs/{job}", headers=headers, json={"company": "Manually corrected"}).status_code
        == 200
    )
    state = client.get("/api/state").json()
    assert state["jobs"][0]["company"] == "Manually corrected"
    assert state["jobs"][0]["original_text"] == original
    assert (
        client.patch(f"/api/jobs/{job}", headers=headers, json={"content_hash": "tamper"}).status_code == 422
    )


def test_profile_and_preference_validation(authenticated):
    client, store, headers = authenticated
    assert client.post("/api/profile", headers=headers, json={"roles": "SQL"}).status_code == 422
    assert client.post("/api/profile", headers=headers, json={"years_experience": -2}).status_code == 422
    assert client.post("/api/preferences", headers=headers, json={"poll_seconds": 1}).status_code == 422
    assert client.post("/api/profile", headers=headers, json={"weights": {"role": 0}}).status_code == 422
    assert (
        client.post("/api/sources", headers=headers, json={"username": "https://evil.example"}).status_code
        == 422
    )


def test_edited_letter_not_automatically_grounded(authenticated):
    client, store, headers = authenticated
    job = store.ingest(1, 1, "Vacancy: Product Analyst\nRequirements: SQL", now())["job_id"]
    profile = store.get("profile")
    profile["facts"] = [
        {
            "id": "approved",
            "text": "I used SQL in a fictional project.",
            "status": "self_reported",
            "approved": True,
            "skills": ["SQL"],
        }
    ]
    store.set("profile", profile)
    letter = client.post(
        "/api/letters", headers=headers, json={"job_id": job, "language": "en", "style": "Professional"}
    ).json()
    r = client.patch(
        f"/api/letters/{letter['id']}",
        headers=headers,
        json={"text": letter["text"] + "\nI increased revenue by 900%."},
    )
    assert r.status_code == 200 and not r.json()["verified"]
    assert r.json()["id"] != letter["id"]
    assert store.rows("SELECT text FROM letters WHERE id=?", (letter["id"],))[0]["text"] == letter["text"]
    assert store.rows("SELECT applied_at FROM applications") == []


def test_archive_pagination_and_original_record_remain_accessible(authenticated):
    client, store, headers = authenticated
    for i in range(51):
        store.ingest(
            1, 1000 + i, f"Vacancy: Product Analyst\nCompany: Fictional {i}\nRequirements: SQL", now()
        )
    first = client.get("/api/jobs?page=1").json()
    second = client.get("/api/jobs?page=2").json()
    assert len(first["items"]) == 50 and len(second["items"]) == 1 and first["total"] == 51
    original = client.get("/api/jobs/" + str(second["items"][0]["id"]))
    assert original.status_code == 200 and "Fictional" in original.json()["original_text"]


def test_backup_restore_and_explicit_data_delete(authenticated):
    client, store, headers = authenticated
    store.ingest(1, 1, "Vacancy: Product Analyst\nRequirements: SQL", now())
    backup = client.get("/api/backup").content
    restored = client.post("/api/restore", headers=headers, files={"file": ("backup.sqlite", backup)})
    assert restored.status_code == 200 and len(store.jobs()) == 1
    assert client.post("/api/delete-data", headers=headers, json={"confirmation": "no"}).status_code == 422
    assert (
        client.post(
            "/api/delete-data", headers=headers, json={"confirmation": "DELETE MY LOCAL DATA"}
        ).status_code
        == 200
    )
    assert store.rows("SELECT * FROM jobs") == []
    assert client.get("/api/state").status_code == 401


def test_html_payload_stays_data_and_public_demo_does_not_expose_private_routes(authenticated):
    client, store, headers = authenticated
    source = store.ingest(
        1, 1, 'Vacancy: Product Analyst\nRequirements: SQL\n<script>alert("xss")</script>', now()
    )
    payload = client.get("/api/state").json()["jobs"][0]["original_text"]
    assert "<script>" in payload
    assert "script-src" in client.get("/").headers["content-security-policy"]
    assert client.get("/private/telegram.session").status_code == 404
    assert client.get("/.env").status_code == 404
    assert source["job_id"]
